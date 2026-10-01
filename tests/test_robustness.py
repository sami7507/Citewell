"""Tests for input validation, error handling and output safety added in the rebuild."""

import io
import time
from unittest.mock import MagicMock, patch

import httpx
import numpy as np
import pytest
from docx import Document as DocxDocument

from citewell import config
from citewell.chunking.chunker import Chunk
from citewell.generation.generator import GenerationError, Generator, _format_source_label
from citewell.ingestion.loaders import DocumentLoadError, load_docx, load_pdf
from citewell.pipeline import chunk_documents
from citewell.ingestion.loaders import PageContent
from citewell.retrieval.retriever import RetrievalResult
from citewell.services.uploads import UploadError, process_uploads, safe_filename
from citewell.storage.feedback import FeedbackEntry, get_all_feedback, log_feedback
from frontend.components import extract_cited_source_numbers, format_answer, friendly_document_name


class FakeUpload:
    def __init__(self, name, content):
        self.name, self._content, self.size = name, content, len(content)

    def getbuffer(self):
        return io.BytesIO(self._content).getbuffer()


# --- config ---

def test_int_setting_falls_back_to_default_on_bad_input(monkeypatch):
    monkeypatch.setenv("TOP_K", "not-a-number")
    config.reload_from_env()
    assert config.TOP_K == 2
    monkeypatch.delenv("TOP_K")
    config.reload_from_env()


# --- loaders ---

def test_corrupt_pdf_raises_friendly_error(tmp_path):
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"%PDF-1.4 this is not really a pdf")
    with pytest.raises(DocumentLoadError, match="broken.pdf"):
        load_pdf(bad)


def test_docx_paragraphs_are_grouped_into_sections_with_headings(tmp_path):
    doc = DocxDocument()
    doc.add_heading("Services", level=1)
    doc.add_paragraph("1. SCOPE. The supplier provides consulting services.")
    doc.add_heading("Payment", level=1)
    doc.add_paragraph("2. FEES. The client pays $5,000 per month.")
    path = tmp_path / "agreement.docx"
    doc.save(path)

    sections = load_docx(path)

    assert len(sections) == 2  # one per heading, not one per paragraph
    assert sections[0].metadata["file_type"] == "docx"
    assert "SCOPE" in sections[0].text and "FEES" in sections[1].text


def test_corrupt_docx_raises_friendly_error(tmp_path):
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"PK\x03\x04 not a real docx")
    with pytest.raises(DocumentLoadError, match="broken.docx"):
        load_docx(bad)


# --- upload validation ---

def test_safe_filename_strips_path_traversal():
    assert safe_filename("../../etc/passwd.pdf") == "passwd.pdf"
    assert safe_filename("C:\\Users\\me\\contract.pdf") == "contract.pdf"
    assert safe_filename("") == "document"


def test_upload_rejects_file_whose_content_does_not_match_extension():
    fake = FakeUpload("report.pdf", b"MZ\x90\x00 this is an executable")
    with pytest.raises(UploadError, match="does not look like a real PDF"):
        process_uploads([fake], MagicMock())


def test_upload_rejects_unsupported_extension_and_empty_batch():
    with pytest.raises(UploadError, match="Only PDF and Word"):
        process_uploads([FakeUpload("notes.txt", b"hello")], MagicMock())
    with pytest.raises(UploadError, match="at least one"):
        process_uploads([], MagicMock())


def test_upload_enforces_file_count_limit(monkeypatch):
    monkeypatch.setattr(config, "MAX_UPLOAD_FILES", 1)
    files = [FakeUpload("a.pdf", b"%PDF-1.4"), FakeUpload("b.pdf", b"%PDF-1.4")]
    with pytest.raises(UploadError, match="at most 1"):
        process_uploads(files, MagicMock())


# --- chunk composition ---

def test_chunk_documents_drops_prose_for_pages_that_have_tables():
    pages = [
        PageContent(text="flattened table junk", source="f.pdf", page_number=1),
        PageContent(text="1. RENT. Pay on time.", source="l.pdf", page_number=1),
    ]
    table = Chunk(text="| a | b |", source="f.pdf", page_number=1, chunk_id="t", chunk_strategy="table", metadata={})

    chunks = chunk_documents(pages, [table])

    assert [c.source for c in chunks if c.chunk_strategy != "table"] == ["l.pdf"]
    assert table in chunks


# --- feedback ---

def test_feedback_rejects_invalid_rating(tmp_path):
    entry = FeedbackEntry(question="q", answer="a", sources=[], rating="meh", mode="m", top_k=2)
    with pytest.raises(ValueError, match="rating"):
        log_feedback(entry, db_path=tmp_path / "f.db")


def test_feedback_truncates_oversized_fields(tmp_path):
    db = tmp_path / "f.db"
    entry = FeedbackEntry(question="q" * 50_000, answer="a", sources=["s"], rating="up", mode="m", top_k=2)
    log_feedback(entry, db_path=db)
    assert len(get_all_feedback(db_path=db)[0]["question"]) == 20_000


# --- answer rendering safety ---

def test_format_answer_escapes_html_so_model_output_cannot_inject_markup():
    out = format_answer('Pay <script>alert(1)</script> now (Source 1).')
    assert "<script>" not in out
    assert "&lt;script&gt;" in out


def test_format_answer_turns_citations_into_numbered_marks_and_escapes_dollars():
    out = format_answer("The rent is $2,400 (Source 1) and the fee is $75 (Source 2).")
    assert out.count('<sup class="cw-ref">') == 2
    assert "\\$2,400" in out
    assert "(Source" not in out


def test_format_answer_handles_grouped_citations():
    out = format_answer("Both apply (Sources 1 and 3).")
    assert out.count('<sup class="cw-ref">') == 2


def test_extract_cited_source_numbers_handles_grouped_citations():
    assert extract_cited_source_numbers("See Sources 1 and 3, and Source 4.") == {1, 3, 4}


def test_friendly_document_name():
    assert friendly_document_name("sample_lease_agreement.pdf") == "Sample lease agreement"


# --- generator error handling and labels ---

def _result(text="Rent is $2,400.", metadata=None, page=1):
    chunk = Chunk(text=text, source="d.docx", page_number=page, chunk_id="x", chunk_strategy="t", metadata=metadata or {})
    return RetrievalResult(chunk=chunk, score=0.9)


def test_docx_citations_say_section_not_page():
    label = _format_source_label(_result(metadata={"file_type": "docx"}, page=3), 1)
    assert label == "Source 1: d.docx, section 3"


def test_authentication_failure_becomes_a_friendly_generation_error():
    import groq

    generator = Generator(api_key="fake")
    response = httpx.Response(401, request=httpx.Request("POST", "https://api.groq.com/x"))
    error = groq.AuthenticationError("bad key", response=response, body=None)

    with patch.object(generator.client.chat.completions, "create", side_effect=error):
        with pytest.raises(GenerationError, match="key was rejected"):
            generator.generate("q", [_result()])


def test_empty_model_answer_is_reported_not_returned():
    generator = Generator(api_key="fake")
    response = MagicMock()
    response.choices = [MagicMock(message=MagicMock(content=None))]

    with patch.object(generator.client.chat.completions, "create", return_value=response):
        with pytest.raises(GenerationError, match="empty answer"):
            generator.generate("q", [_result()])


def test_oversized_chunk_text_is_truncated_in_the_prompt(monkeypatch):
    monkeypatch.setattr(config, "MAX_CONTEXT_CHARS_PER_CHUNK", 100)
    generator = Generator(api_key="fake")
    response = MagicMock()
    response.choices = [MagicMock(message=MagicMock(content="ok (Source 1)"))]

    with patch.object(generator.client.chat.completions, "create", return_value=response) as create:
        generator.generate("q", [_result(text="word " * 500)])

    assert "[excerpt truncated]" in create.call_args.kwargs["messages"][1]["content"]
